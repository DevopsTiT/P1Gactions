#!/usr/bin/env python3
"""Architecture map PPT for Dynatrace ↔ ServiceNow."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
NAVY = RGBColor(0x0B, 0x2C, 0x5C)
NAVY_MID = RGBColor(0x14, 0x3D, 0x7A)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREEN = RGBColor(0x1E, 0x84, 0x4E)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
LIGHT_BLUE = RGBColor(0xD6, 0xE6, 0xF5)
LIGHT_GREEN = RGBColor(0xD5, 0xF5, 0xE3)
LIGHT_ORANGE = RGBColor(0xFD, 0xE8, 0xD0)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, text, size=12, bold=False, color=BLACK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_textbox(slide, left, top, width, height, text, size=12, bold=False, color=BLACK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)


def add_bullets(slide, left, top, width, height, lines, size=14):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        set_run(run, line, size=size, color=BLACK)


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def header_bar(slide, title, subtitle=""):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.95), NAVY)
    add_textbox(slide, Inches(0.4), Inches(0.18), Inches(12), Inches(0.45), title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.4), Inches(0.55), Inches(12), Inches(0.35), subtitle, size=12, color=LIGHT_BLUE)


def card(slide, left, top, width, height, title, body, fill=WHITE, title_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = NAVY_MID
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_top = Inches(0.08)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, title, size=13, bold=True, color=title_color)
    if body:
        for line in body.split("\n"):
            p = tf.add_paragraph()
            r = p.add_run()
            set_run(r, line, size=11, color=GRAY)
    return shape


def arrow_label(slide, left, top, text):
    add_textbox(slide, left, top, Inches(1.2), Inches(0.35), text, size=11, bold=True, color=TEAL)


def build(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, NAVY)
    add_textbox(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(1), "Dynatrace ↔ ServiceNow Architecture", size=32, bold=True, color=WHITE)
    add_textbox(s, Inches(0.8), Inches(3.3), Inches(11.5), Inches(0.6), "System context · Incident path · Three lanes · Control points", size=18, color=LIGHT_BLUE)

    # 2 system context
    s = prs.slides.add_slide(blank)
    header_bar(s, "System context", "Operators sit above both systems")
    card(s, Inches(4.5), Inches(1.15), Inches(4.3), Inches(1.1), "On-call / Operators", "ServiceNow UI + Dynatrace UI", LIGHT_ORANGE)
    card(s, Inches(0.4), Inches(2.6), Inches(5.8), Inches(3.5), "Dynatrace", "OneAgent / RUM\nSmartscape + Davis\nProblems\nAlerting profiles\nProblem notifications", LIGHT_BLUE, TEAL)
    card(s, Inches(7.1), Inches(2.6), Inches(5.8), Inches(3.5), "ServiceNow", "ITSM Incidents\nTransform + Assignment\nCMDB / Service Graph CIs\nOptional ITOM Events", LIGHT_GREEN, GREEN)
    add_textbox(s, Inches(0.4), Inches(6.3), Inches(12), Inches(0.6), "Monitored apps/hosts feed Dynatrace; optional CMDB sync makes the same entities into ServiceNow CIs.", size=13, color=GRAY)

    # 3 incident path
    s = prs.slides.add_slide(blank)
    header_bar(s, "Incident path (main architecture)", "Left → right: detect to resolve")
    card(s, Inches(0.2), Inches(1.3), Inches(2.0), Inches(2.2), "1 App/Host", "eip-checkout\nOneAgent", LIGHT_GRAY)
    card(s, Inches(2.4), Inches(1.3), Inches(2.0), Inches(2.2), "2 Davis", "Problem\n+ RCA", LIGHT_BLUE)
    card(s, Inches(4.6), Inches(1.3), Inches(2.0), Inches(2.2), "3 Profile", "prod filter\nyes / no", LIGHT_ORANGE)
    card(s, Inches(6.8), Inches(1.3), Inches(2.0), Inches(2.2), "4 Notify", "ServiceNow\nITSM ON", LIGHT_BLUE)
    card(s, Inches(9.0), Inches(1.3), Inches(2.0), Inches(2.2), "5 Transform", "Import set\n→ fields", LIGHT_GREEN)
    card(s, Inches(11.1), Inches(1.3), Inches(2.0), Inches(2.2), "6 Incident", "CI + group\n+ Problem URL", LIGHT_GREEN)
    add_bullets(s, Inches(0.4), Inches(3.8), Inches(12), Inches(3), [
        "No profile match → stop (no ticket)",
        "On-call opens Problem URL from Incident → fixes app/infra",
        "Problem CLOSED → notification update → Incident Resolved",
    ], size=16)

    # 4 three lanes
    s = prs.slides.add_slide(blank)
    header_bar(s, "Three lanes", "Same Dynatrace source, different ServiceNow targets")
    card(s, Inches(0.3), Inches(1.3), Inches(4.0), Inches(5.0), "A. Incident (default)", "Direction: DT → SNOW\n\nProblem open/update/close\n→ Incident\n\nDay-2 SRE tickets", LIGHT_BLUE, TEAL)
    card(s, Inches(4.6), Inches(1.3), Inches(4.0), Inches(5.0), "B. Events (optional)", "Direction: DT → SNOW\n\nEvents → em_event\nITOM rules decide tickets\n\nNeeds Event Mgmt", LIGHT_ORANGE, ORANGE)
    card(s, Inches(8.9), Inches(1.3), Inches(4.0), Inches(5.0), "C. CMDB / Service Graph", "Direction: DT → SNOW pull\n\nHosts/services/apps → CIs\n\nImproves Affected CI", LIGHT_GREEN, GREEN)

    # 5 control points
    s = prs.slides.add_slide(blank)
    header_bar(s, "Control points (where manage lives)", "Eight knobs across the architecture")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(5.2), "Dynatrace", "1 Entity tags (env, app)\n2 Alerting profile (noise)\n3 Problem notification (wire)\n\nHealthy: STG never hits prod queue\nOne notify per environment", LIGHT_BLUE)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), "ServiceNow", "4 Integration user (least privilege)\n5 Transform map (fields)\n6 Assignment rules (group)\n7 CI identity (dedupe)\n8 Import set health (errors)\n\nHealthy: CI filled, right group, close sync", LIGHT_GREEN)

    # 6 EIP overlay
    s = prs.slides.add_slide(blank)
    header_bar(s, "EIP example on the same map", "Filled architecture for INC0012345")
    add_bullets(s, Inches(0.5), Inches(1.3), Inches(12), Inches(5.5), [
        "Shopper → eip-checkout slow",
        "Davis Problem P-240906 (DB behind eip-checkout)",
        "Profile prod-eip-to-servicenow MATCHES",
        "Notify → Transform → INC0012345",
        "CI: eip-checkout / eip-app-01 | Group: EIP-Support | Notes: Problem URL",
        "On-call fix → Problem CLOSED → Incident Resolved",
    ], size=17)

    # 7 trust boundary
    s = prs.slides.add_slide(blank)
    header_bar(s, "Trust boundary", "What crosses the wire")
    card(s, Inches(0.3), Inches(1.3), Inches(6.2), Inches(4.5), "Dynatrace tenant", "Problems, entities, RCA\nEgress IPs must be allowed\nNo real PII in test payloads", LIGHT_BLUE)
    card(s, Inches(6.8), Inches(1.3), Inches(6.1), Inches(4.5), "ServiceNow instance", "Incidents, CMDB, users\nIntegration user scoped\nAudit on ticket changes\nAuth / IP failures → 403", LIGHT_GREEN)
    add_textbox(s, Inches(0.5), Inches(6.1), Inches(12), Inches(0.7), "Wire carries metadata (title, severity, URL, entities, state) — not a full log dump.", size=14, bold=True, color=NAVY)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"wrote {out_path} slides {len(prs.slides)}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    build(here / "27-ServiceNow-Dynatrace-Architecture.pptx")
