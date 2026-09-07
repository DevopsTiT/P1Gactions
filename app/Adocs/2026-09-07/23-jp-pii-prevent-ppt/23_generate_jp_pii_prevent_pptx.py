#!/usr/bin/env python3
"""JP PII prevent PPT — OneAgent + OpenPipeline + verify (no app layer)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x2C, 0x5C)
NAVY_MID = RGBColor(0x14, 0x3D, 0x7A)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREEN = RGBColor(0x1E, 0x84, 0x4E)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
ACCENT_RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT_BLUE = RGBColor(0xD6, 0xE6, 0xF5)
LIGHT_GREEN = RGBColor(0xD5, 0xF5, 0xE3)
LIGHT_ORANGE = RGBColor(0xFD, 0xE8, 0xD0)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
LIGHT_RED = RGBColor(0xFA, 0xE5, 0xD3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, text, size=12, bold=False, color=BLACK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_textbox(slide, left, top, width, height, text, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, lines, size=14, color=BLACK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        set_run(run, line, size=size, color=color)
    return box


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def header_bar(slide, title, subtitle=""):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.95), NAVY)
    add_textbox(slide, Inches(0.4), Inches(0.18), Inches(12), Inches(0.45), title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.4), Inches(0.55), Inches(12), Inches(0.35), subtitle, size=12, color=LIGHT_BLUE)


def card(slide, left, top, width, height, title, body, fill=WHITE, title_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = NAVY_MID
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.1)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, title, size=14, bold=True, color=title_color)
    if body:
        for line in body.split("\n"):
            p = tf.add_paragraph()
            r = p.add_run()
            set_run(r, line, size=11, color=GRAY)
    return shape


def build(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, NAVY)
    add_textbox(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1), "Prevent JP PII in Dynatrace Logs", size=36, bold=True, color=WHITE)
    add_textbox(s, Inches(0.8), Inches(3.2), Inches(11.5), Inches(0.6), "OneAgent masking + OpenPipeline + DQL verify (no app-layer scrub)", size=18, color=LIGHT_BLUE)
    add_textbox(s, Inches(0.8), Inches(4.2), Inches(11.5), Inches(0.5), "Basis: Splunk JP-chars → PII keyword filter (氏名 / 住所 / 電話 / …)", size=14, color=WHITE)

    # 2 agenda
    s = prs.slides.add_slide(blank)
    header_bar(s, "Agenda", "What this deck covers")
    add_bullets(s, Inches(0.6), Inches(1.3), Inches(12), Inches(5.5), [
        "1. Goal: stop raw JP PII landing in Grail",
        "2. Basis from Splunk: Japanese text first, then keyword list",
        "3. Step A — OneAgent Sensitive data masking",
        "4. Step B — OpenPipeline content mask",
        "5. Step C — DQL verify",
        "6. PII examples + filter / mask patterns",
        "7. Checklist",
    ], size=18)

    # 3 goal
    s = prs.slides.add_slide(blank)
    header_bar(s, "Goal", "Prevent future ingest — not only scan after the fact")
    card(s, Inches(0.5), Inches(1.3), Inches(4), Inches(2.2), "Find (Splunk/DQL)", "JP chars + keywords\nShows where PII already is\nDoes NOT block future logs", LIGHT_ORANGE, ORANGE)
    card(s, Inches(4.7), Inches(1.3), Inches(4), Inches(2.2), "Prevent (Dynatrace)", "OneAgent mask → ***\nOpenPipeline mask → ***\nBefore / at ingest", LIGHT_GREEN, GREEN)
    card(s, Inches(8.9), Inches(1.3), Inches(3.9), Inches(2.2), "Verify", "Same keyword idea in DQL\nExpect 0 raw hits or only ***", LIGHT_BLUE, NAVY)
    add_textbox(s, Inches(0.5), Inches(3.9), Inches(12), Inches(2.5),
                "Flow:  Fake/test log with JP PII labels  →  OneAgent  →  OpenPipeline  →  Grail (masked)  →  DQL check",
                size=16, bold=True, color=NAVY)

    # 4 basis
    s = prs.slides.add_slide(blank)
    header_bar(s, "Basis — same as Splunk chat", "Two filters in order")
    card(s, Inches(0.5), Inches(1.3), Inches(6), Inches(2.8), "1) Japanese characters",
         "Unicode ranges\nHiragana/Katakana/Kanji\n[\\u3000-\\u30FF\\u4E00-\\u9FFF]\nCuts English-only noise first", LIGHT_BLUE)
    card(s, Inches(6.8), Inches(1.3), Inches(6), Inches(2.8), "2) JP PII keywords",
         "氏名 住所 電話 生年月日\nメール マイナンバー 口座番号\nクレジットカード …\nThen mask values next to labels", LIGHT_GREEN)
    add_textbox(s, Inches(0.5), Inches(4.4), Inches(12), Inches(2),
                "Dynatrace: use keyword list for OneAgent/OpenPipeline regex. Use JP+keyword DQL only to verify.",
                size=15, color=GRAY)

    # 5 steps overview
    s = prs.slides.add_slide(blank)
    header_bar(s, "Three steps (no app layer)", "A → B → C")
    card(s, Inches(0.4), Inches(1.4), Inches(4), Inches(4.5), "A. OneAgent",
         "Settings → Log monitoring\n→ Configure log module\n→ Sensitive data masking\n\nSTRING → ***\nHost-group scope\nWave 1 keywords first\nFake-data test", LIGHT_GREEN, GREEN)
    card(s, Inches(4.6), Inches(1.4), Inches(4), Inches(4.5), "B. OpenPipeline",
         "OpenPipeline → Logs\nContent mask same regex\n\nSecond net if OneAgent missed\nPrefer mask over drop\nFake-data test again", LIGHT_BLUE, NAVY)
    card(s, Inches(8.8), Inches(1.4), Inches(4), Inches(4.5), "C. DQL verify",
         "Logs / Notebooks\nJP chars then keywords\n\n0 hits or *** only = pass\nRaw values = fix A/B\nWeekly re-check", LIGHT_ORANGE, ORANGE)

    # 6 OneAgent detail
    s = prs.slides.add_slide(blank)
    header_bar(s, "Step A — OneAgent UI path", "Sensitive data masking")
    add_bullets(s, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5), [
        "1. Settings → Collect and capture → Log monitoring",
        "2. Configure log module → Sensitive data masking",
        "3. Scope: Host group (safer) or Environment",
        "4. Each rule: Masking type STRING | Replacement ***",
        "5. Search expression: regex; capture VALUE after label",
        "6. Add Wave 1 rules one-by-one → Save → fake test",
        "7. Expand Wave 2 / Wave 3 after proof",
        "UI: enable generic email mask if available",
    ], size=16)

    # 7 OpenPipeline
    s = prs.slides.add_slide(blank)
    header_bar(s, "Step B — OpenPipeline", "Second net at ingest")
    add_bullets(s, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5), [
        "1. Open OpenPipeline → Logs pipeline for your env",
        "2. Add content mask / sensitive-data processor",
        "3. Reuse same Wave 1 JP keyword patterns → ***",
        "4. Prefer MASK (keep ERROR logs for ops)",
        "5. Drop entire record only if privacy requires",
        "6. Fake-test on the OpenPipeline ingest path",
        "7. Then expand keywords",
    ], size=16)

    # 8 verify
    s = prs.slides.add_slide(blank)
    header_bar(s, "Step C — DQL verify", "Same order as Splunk")
    add_textbox(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(3.2),
                "fetch logs\n| filter matchesValue(content, \".*[\\\\u3000-\\\\u30FF\\\\u4E00-\\\\u9FFF].*\")\n| filter matchesValue(content, \".*(氏名|住所|電話|メール|マイナンバー|口座番号|クレジットカード).*\")\n| summarize hits = count(), by: { host.name }\n| sort hits desc",
                size=14, color=BLACK)
    add_bullets(s, Inches(0.5), Inches(4.6), Inches(12), Inches(2.2), [
        "0 hits → good for that timeframe / keywords",
        "Hits with *** only → mask working",
        "Hits with real names/numbers → fix OneAgent/OpenPipeline regex or scope",
    ], size=15)

    # 9 PII categories examples
    s = prs.slides.add_slide(blank)
    header_bar(s, "PII keyword categories (examples)", "Filter / mask like the Splunk list")
    card(s, Inches(0.3), Inches(1.2), Inches(4.1), Inches(2.5), "Identity",
         "氏名 年齢 性別\n生年月日 誕生日\n生年 生月 生日 生年月", LIGHT_BLUE)
    card(s, Inches(4.6), Inches(1.2), Inches(4.1), Inches(2.5), "Contact",
         "住所 郵便番号\n電話 電話番号 固定電話\n携帯 携帯電話 携帯番号\nメール メールアドレス ファックス", LIGHT_GREEN)
    card(s, Inches(8.9), Inches(1.2), Inches(4), Inches(2.5), "Government ID",
         "身分証 身分証明書\nマイナンバー 個人番号\n運転免許 免許証\nパスポート 旅券", LIGHT_ORANGE)
    card(s, Inches(0.3), Inches(4.0), Inches(6.2), Inches(2.6), "Financial",
         "銀行口座 口座番号 銀行コード 支店コード\nクレジットカード カード番号 カード", LIGHT_RED)
    card(s, Inches(6.7), Inches(4.0), Inches(6.2), Inches(2.6), "Other / noisy",
         "会社 — broad, false positives likely\nカード alone — noisy; prefer カード番号", LIGHT_GRAY)

    # 10 mask pattern examples
    s = prs.slides.add_slide(blank)
    header_bar(s, "How to filter / mask — pattern examples", "OneAgent & OpenPipeline search expressions")
    add_bullets(s, Inches(0.4), Inches(1.15), Inches(12.5), Inches(5.8), [
        "氏名\\s*[:：＝=]\\s*(\\S+)     → mask name value",
        "住所\\s*[:：＝=]\\s*(.+)       → mask address value",
        "(電話番号?|携帯(電話|番号)?)\\s*[:：＝=]\\s*(\\S+)",
        "(メール(アドレス)?)\\s*[:：＝=]\\s*(\\S+)",
        "(マイナンバー|個人番号)\\s*[:：＝=]\\s*(\\S+)",
        "(口座番号|銀行口座)\\s*[:：＝=]\\s*(\\S+)",
        "(クレジットカード|カード番号)\\s*[:：＝=]\\s*(\\S+)",
        'JSON later: "bankAccountNo"\\s*:\\s*"(.*?)"',
        "Fake test: 氏名:テスト太郎 口座番号:0000000 → expect ***",
    ], size=14)

    # 11 waves
    s = prs.slides.add_slide(blank)
    header_bar(s, "Roll out in waves", "Small blast radius")
    card(s, Inches(0.4), Inches(1.3), Inches(4), Inches(4.5), "Wave 1",
         "氏名\n住所\n電話 / 電話番号\nメール\nマイナンバー\n口座番号", LIGHT_GREEN, GREEN)
    card(s, Inches(4.6), Inches(1.3), Inches(4), Inches(4.5), "Wave 2",
         "生年月日 / 郵便番号\n携帯*\n身分証*\nパスポート\nクレジットカード", LIGHT_BLUE, NAVY)
    card(s, Inches(8.8), Inches(1.3), Inches(4), Inches(4.5), "Wave 3",
         "Rest of keyword list\n+ English JSON keys\nbankAccountNo\npolicyHolderName\nphones / emails", LIGHT_ORANGE, ORANGE)

    # 12 checklist
    s = prs.slides.add_slide(blank)
    header_bar(s, "Checklist", "Pass criteria")
    add_bullets(s, Inches(0.5), Inches(1.3), Inches(12), Inches(5.5), [
        "A — OneAgent Wave 1 on host group → fake line shows ***",
        "B — OpenPipeline same patterns → fake line still safe",
        "C — DQL verify → 0 raw PII (or *** only)",
        "Expand Wave 2–3 only after Wave 1 passes",
        "Weekly DQL re-check for regressions",
        "Never test with real customer PII",
        "Do not drop all Japanese logs — mask values",
    ], size=17)

    # 13 closing
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, NAVY)
    add_textbox(s, Inches(0.8), Inches(2.5), Inches(11.5), Inches(1), "OneAgent → OpenPipeline → Verify", size=32, bold=True, color=WHITE)
    add_textbox(s, Inches(0.8), Inches(3.6), Inches(11.5), Inches(1), "Same JP PII keyword basis as Splunk — Dynatrace masks before Grail", size=16, color=LIGHT_BLUE)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print("wrote", out_path, "slides", len(prs.slides))


if __name__ == "__main__":
    base = Path("/Users/k/Codes/Pra/P1GithubActions/P1Gactions/app/Adocs/2026-09-07/23-jp-pii-prevent-ppt")
    build(base / "23-JP-PII-Prevent-OneAgent-OpenPipeline.pptx")
