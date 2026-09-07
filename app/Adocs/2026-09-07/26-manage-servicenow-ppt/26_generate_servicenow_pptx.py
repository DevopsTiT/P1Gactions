#!/usr/bin/env python3
"""ServiceNow + Dynatrace manage PPT — setup, day-2, EIP example."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x2C, 0x5C)
NAVY_MID = RGBColor(0x14, 0x3D, 0x7A)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREEN = RGBColor(0x1E, 0x84, 0x4E)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
LIGHT_BLUE = RGBColor(0xD6, 0xE6, 0xF5)
LIGHT_GREEN = RGBColor(0xD5, 0xF5, 0xE3)
LIGHT_ORANGE = RGBColor(0xFD, 0xE8, 0xD0)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, text, size=12, bold=False, color=BLACK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_textbox(slide, left, top, width, height, text, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, lines, size=14, color=BLACK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        set_run(run, line, size=size, color=color)
    return box


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def header_bar(slide, title, subtitle=""):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.95), NAVY)
    add_textbox(slide, Inches(0.4), Inches(0.18), Inches(12), Inches(0.45), title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.4), Inches(0.55), Inches(12), Inches(0.35), subtitle, size=12, color=LIGHT_BLUE)


def card(slide, left, top, width, height, title, body, fill=WHITE, title_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = NAVY_MID
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.1)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, title, size=14, bold=True, color=title_color)
    if body:
        for line in body.split("\n"):
            p = tf.add_paragraph()
            r = p.add_run()
            set_run(r, line, size=11, color=GRAY)
    return shape


def build(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, NAVY)
    add_textbox(s, Inches(0.8), Inches(2.1), Inches(11.5), Inches(1), "Manage ServiceNow with Dynatrace", size=34, bold=True, color=WHITE)
    add_textbox(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(0.6), "Setup steps + day-2 manage + EIP checkout example", size=18, color=LIGHT_BLUE)
    add_textbox(s, Inches(0.8), Inches(4.1), Inches(11.5), Inches(0.5), "Problem → Incident → right CI / group → fix → both close", size=14, color=WHITE)

    # 2 agenda
    s = prs.slides.add_slide(blank)
    header_bar(s, "Agenda", "What this deck covers")
    add_bullets(s, Inches(0.6), Inches(1.3), Inches(12), Inches(5.5), [
        "1. What ServiceNow + Dynatrace each do",
        "2. Three integration lanes (Incident / Events / CMDB)",
        "3. First-time setup (7 steps)",
        "4. Day-2 manage checklist",
        "5. Detailed example: EIP checkout → INC0012345",
        "6. On-call ticket handling + troubleshooting",
        "7. Hard rules (no ticket storms)",
    ], size=18)

    # 3 goal / pieces
    s = prs.slides.add_slide(blank)
    header_bar(s, "What you are managing", "Tickets + CMDB ↔ Problems + RCA")
    card(s, Inches(0.3), Inches(1.2), Inches(4.0), Inches(2.4), "Dynatrace", "Detects Problems\nDavis root cause\nProblem URL = truth", LIGHT_BLUE)
    card(s, Inches(4.6), Inches(1.2), Inches(4.0), Inches(2.4), "ServiceNow", "Incidents (work tickets)\nCMDB CIs (inventory)\nAssignment + SLA", LIGHT_GREEN)
    card(s, Inches(8.9), Inches(1.2), Inches(4.0), Inches(2.4), "Manage means", "Who gets tickets\nWhich alerts fire\nCI accuracy\nOpen / close loop", LIGHT_ORANGE)
    add_bullets(s, Inches(0.5), Inches(3.9), Inches(12), Inches(3), [
        "Incident = work ticket for humans (ownership, SLA, audit)",
        "CMDB CI = host / service / app the ticket hangs on",
        "Alerting profile = Dynatrace filter (stops STG noise)",
        "Problem notification = the wire that creates / updates the Incident",
    ], size=15)

    # 4 three lanes
    s = prs.slides.add_slide(blank)
    header_bar(s, "Three integration lanes", "Start with Incident; add CMDB when ready")
    card(s, Inches(0.3), Inches(1.3), Inches(4.0), Inches(4.8), "1. Incident (start here)", "ITSM + Dynatrace Incident Integration\n\nEach Problem → Incident\n\nOwner: SRE + ITSM admin", LIGHT_BLUE, TEAL)
    card(s, Inches(4.6), Inches(1.3), Inches(4.0), Inches(4.8), "2. Events (optional)", "ServiceNow ITOM\n\nEvents → em_event\nRules decide tickets\n\nOwner: Event Mgmt", LIGHT_ORANGE, ORANGE)
    card(s, Inches(8.9), Inches(1.3), Inches(4.0), Inches(4.8), "3. CMDB / Service Graph", "Hosts / services → CIs\n\nBetter Affected CI link\n\nOwner: CMDB + platform", LIGHT_GREEN, GREEN)

    # 5 flow
    s = prs.slides.add_slide(blank)
    header_bar(s, "Happy path flow", "Detect → filter → ticket → fix → close")
    add_bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.5), [
        "1. App slows / errors",
        "2. Dynatrace Davis opens a Problem (+ root cause entity)",
        "3. Alerting profile allows notify (prod + severity + tags)",
        "4. Problem notification → ServiceNow",
        "5. Import set → Transform map → Incident",
        "6. Affected CI + assignment group filled",
        "7. On-call uses Problem URL for RCA → fixes",
        "8. Problem CLOSED → Incident Resolved",
    ], size=17)

    # 6 setup ServiceNow
    s = prs.slides.add_slide(blank)
    header_bar(s, "Setup Steps 1–3 — ServiceNow first", "Install before Dynatrace notify")
    card(s, Inches(0.3), Inches(1.2), Inches(4.0), Inches(5.2), "Step 1 — Install app", "ServiceNow Store\nDynatrace Incident Integration\nRun Guided Setup\n\nPass: app is active", LIGHT_BLUE)
    card(s, Inches(4.6), Inches(1.2), Inches(4.0), Inches(5.2), "Step 2 — Integration user", "User: dynatrace.integration\nLeast-privilege roles\nStrong password / OAuth\n\nURL: https://id.service-now.com", LIGHT_GREEN)
    card(s, Inches(8.9), Inches(1.2), Inches(4.0), Inches(5.2), "Step 3 — Transform + assign", "Title → short_description\nSeverity → priority\nURL → work notes\nEntity → Affected CI\n\nEIP apps → EIP-Support", LIGHT_ORANGE)

    # 7 setup Dynatrace
    s = prs.slides.add_slide(blank)
    header_bar(s, "Setup Steps 4–5 — Dynatrace", "Noise filter then notification")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(5.2), "Step 4 — Alerting profile", "Name: prod-eip-to-servicenow\n\nProduction management zone\nSeverity: Error + Critical\nTags: env:prod (pilot: app:eip)\n\nPass: STG does NOT match", LIGHT_ORANGE)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), "Step 5 — Problem notification", "Settings → Integration\n→ Problem notifications\n→ type ServiceNow\n\nSend incidents into ITSM = ON\nAttach alerting profile\nSend test → Save\n\nPass: test Incident appears", LIGHT_BLUE)

    # 8 CMDB + closed loop
    s = prs.slides.add_slide(blank)
    header_bar(s, "Setup Steps 6–7 — CMDB + closed loop", "Optional inventory, then prove resolve sync")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(5.2), "Step 6 — CMDB (optional)", "Service Graph Connector\nfor Observability – Dynatrace\n\nAPI token (read entities)\nSchedule host/service sync\nIdentity: merge by hostname\n\nPass: Affected CI usually filled", LIGHT_GREEN)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), "Step 7 — Closed-loop test", "Open test / fake Problem\n→ Incident New / In Progress\n\nClose Dynatrace Problem\n→ Incident Resolved\n\nPass: no zombie open tickets", LIGHT_ORANGE)

    # 9 day-2
    s = prs.slides.add_slide(blank)
    header_bar(s, "Day-2 manage", "Weekly hygiene on both sides")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(5.2), "Dynatrace weekly", "1. Re-send test notification\n2. Review alerting profile\n3. Open Problems ≈ open DT Incidents\n4. Problem URL still in notes\n\nHealthy: no STG flood", LIGHT_BLUE)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), "ServiceNow weekly", "1. Import set errors = 0\n2. Right assignment group\n3. CI link rate high\n4. One Problem → one Incident\n\nHealthy: updates, not clones", LIGHT_GREEN)

    # 10 hard rules
    s = prs.slides.add_slide(blank)
    header_bar(s, "Hard rules", "Do not break these")
    add_bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.5), [
        "Do NOT send all Problems from all environments → ticket storm",
        "Prefer one Problem notification per environment",
        "Tag entities (env, app) for filter + assignment + CMDB match",
        "Keep Problem URL as source of truth for RCA",
        "Never put real customer PII in test tickets",
        "On-call opens Problem URL — ticket text alone is not enough",
    ], size=18)

    # 11 example story
    s = prs.slides.add_slide(blank)
    header_bar(s, "Example story — EIP checkout", "One Problem → one Incident → fix → both close")
    add_textbox(s, Inches(0.5), Inches(1.2), Inches(12), Inches(0.5), "Shoppers see slow checkout. You want one ticket with correct CI and Davis RCA.", size=16, bold=True, color=NAVY)
    add_bullets(s, Inches(0.6), Inches(1.9), Inches(12), Inches(5), [
        "Actors: Shopper | Dynatrace Davis | ServiceNow | On-call SRE | App/DBA",
        "Goal: INC for EIP-Support — not five Slack threads",
        "Tools on-call uses: Problem URL + multi-app dashboard (app=EIP)",
        "Outcome: rollback / DB fix → Problem CLOSED → Incident Resolved",
    ], size=17)

    # 12 timeline
    s = prs.slides.add_slide(blank)
    header_bar(s, "Example timeline", "EIP checkout → INC0012345")
    add_bullets(s, Inches(0.5), Inches(1.25), Inches(12.2), Inches(5.8), [
        "T+0   Checkout P95 up; errors rise",
        "T+2m  Problem P-240906 — Davis: slow DB behind eip-checkout",
        "T+2m  Profile prod-eip-to-servicenow matches → notify",
        "T+3m  Import → Transform → Incident INC0012345",
        "T+3m  CI = eip-checkout / eip-app-01 | Group = EIP-Support",
        "T+5m  On-call acks → Problem URL → dashboard app=EIP",
        "T+20m Rollback / DB fix",
        "T+25m Dynatrace closes Problem",
        "T+26m Incident → Resolved",
    ], size=15)

    # 13 INC fields + on-call
    s = prs.slides.add_slide(blank)
    header_bar(s, "What INC0012345 looks like + on-call steps", "Ticket content and human workflow")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(5.2), "Incident fields", "Number: INC0012345\nShort desc: [Dynatrace] Checkout latency — eip-checkout\nNotes: Davis summary + Problem URL\nPriority: from severity map\nCI: eip-checkout / eip-app-01\nGroup: EIP-Support\nOpened by: dynatrace.integration", LIGHT_BLUE)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), "On-call steps", "1. Ack Incident\n2. Open Problem URL\n3. Read Davis root cause\n4. Dashboard filter app=EIP\n5. Fix or escalate + work notes\n6. Confirm Problem close → Resolved\n7. Recurring? → Problem Mgmt", LIGHT_GREEN)

    # 14 troubleshooting
    s = prs.slides.add_slide(blank)
    header_bar(s, "Troubleshooting", "When manage is broken")
    add_bullets(s, Inches(0.5), Inches(1.3), Inches(12.2), Inches(5.8), [
        "Test notify 403 → auth / IP allow list for Dynatrace → ServiceNow",
        "No Incident → ITSM flag off, or transform / import set errors",
        "Duplicate Incidents → multiple notifications; use update-on-change",
        "Blank Affected CI → CMDB not synced / hostname mismatch",
        "Wrong group → assignment rule / tag mapping",
        "Incident stays open → close state not mapped from Problem CLOSED",
        "Ticket storm → alerting profile too wide (add MZ + severity + env)",
    ], size=15)

    # 15 before/after + close
    s = prs.slides.add_slide(blank)
    header_bar(s, "Before vs after + takeaway", "What good manage looks like")
    card(s, Inches(0.3), Inches(1.2), Inches(6.2), Inches(3.8), "Before", "Slack flood, no owner\nGuess which server\nFix done, ticket forgotten\nSTG wakes prod on-call", LIGHT_ORANGE)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(3.8), "After", "One Incident with SLA\nCI already on ticket\nProblem close → Resolved\nProfile blocks STG", LIGHT_GREEN)
    add_textbox(s, Inches(0.5), Inches(5.3), Inches(12), Inches(1.2),
                "Manage = notification + alerting profile + transform/assignment + CI match + weekly hygiene.\nExample = EIP Problem → INC0012345 → RCA via Problem URL → both close.",
                size=15, bold=True, color=NAVY)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"wrote {out_path} slides {len(prs.slides)}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    build(here / "26-Manage-ServiceNow-Dynatrace.pptx")
