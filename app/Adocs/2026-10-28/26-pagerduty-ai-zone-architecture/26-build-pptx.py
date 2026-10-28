#!/usr/bin/env python3
"""PagerDuty AI architecture in the zone style of the APM architecture page."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = (
    "/Users/k/Learnings/AIProject/CursorFiles/Daily Files/"
    "2026-10-28/26-pagerduty-ai-zone-architecture/"
    "26-pagerduty-ai-zone-architecture.pptx"
)

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY2 = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x0E, 0x6B, 0x5C)
TEAL_BG = RGBColor(0xE7, 0xF6, 0xF2)
BLUE_BG = RGBColor(0xE8, 0xF1, 0xFA)
ORANGE = RGBColor(0xC4, 0x56, 0x0A)
ORANGE_BG = RGBColor(0xFF, 0xF4, 0xE8)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
GREEN_BG = RGBColor(0xE8, 0xF6, 0xEE)
SLATE = RGBColor(0x37, 0x41, 0x51)
SLATE_BG = RGBColor(0xF4, 0xF6, 0xF8)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xF6, 0xCC)
GOLD_LINE = RGBColor(0xC4, 0x96, 0x0C)
BG = RGBColor(0xF7, 0xF8, 0xFB)
ZONE = RGBColor(0xEE, 0xF3, 0xF8)


def set_run(paragraph, text, size, bold, color):
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.01)
    first = True
    for text, size, bold, color, before in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(before)
        p.space_after = Pt(0)
        set_run(p, text, size, bold, color)
    return box


def add_card(slide, x, y, w, h, fill, border, border_pt=1.15, roundness=0.08):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = roundness
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    shape.shadow.inherit = False
    return shape


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = NAVY
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()
    head = etree.SubElement(ln, qn("a:headEnd"))
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("h", "med")


def red_dot(slide, x, y):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.14), Inches(0.14))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RED
    dot.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "PagerDuty AI Architecture"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.56))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, 0.26, 0.04, 8.5, 0.28, [("PagerDuty AI Architecture", 22, True, WHITE, 0)])
    add_text(slide, 0.28, 0.3, 9.2, 0.22, [("Zone picture for the 24x7 war-room assist. ALJ and ADJ (P&C).", 12, False, RGBColor(0xD6, 0xE4, 0xF0), 0)])

    add_card(slide, 0.18, 0.66, 13.0, 0.48, GOLD, GOLD_LINE, 1.0)
    add_text(slide, 0.28, 0.7, 12.8, 0.4, [
        ("Same kind of picture as the APM architecture page: zones, a platform in the middle, a label on every line.", 12, False, SLATE, 0),
        ("The only new box is PagerDuty AI, inside the PagerDuty cloud. This is not the APM or MultiSoft platform.", 12, False, SLATE, 0),
    ])

    # Zone 1
    add_card(slide, 0.18, 1.26, 3.15, 4.45, ZONE, NAVY2, 1.35, 0.04)
    add_text(slide, 0.28, 1.3, 2.95, 0.28, [("1  Signal sources", 13, True, NAVY, 0)])
    sources = [
        (1.64, "Dynatrace", "APM, user, synthetic, host"),
        (2.28, "Splunk", "Logs and log alerts"),
        (2.92, "Selenium / Jenkins", "Scripted checks"),
        (3.56, "Network and VM", "VeloCloud, Nagios, SCOM"),
        (4.2, "Cloud tools", "CloudWatch, X-Ray, Prometheus"),
    ]
    for y, title, sub in sources:
        add_card(slide, 0.32, y, 2.88, 0.56, WHITE, NAVY2, 1.0)
        add_text(slide, 0.4, y + 0.02, 2.72, 0.5, [(title, 12, True, NAVY, 0), (sub, 11, False, SLATE, 0)])
    add_text(slide, 0.32, 4.82, 2.88, 0.75, [
        ("Each tool keeps its current owner.", 11, False, SLATE, 0),
        ("Technology DevOps owns Dynatrace, Splunk, and Jenkins.", 11, False, SLATE, 1),
    ])

    # Zone 2 PagerDuty cloud
    add_card(slide, 4.05, 1.26, 5.15, 4.45, RGBColor(0xE8, 0xEE, 0xF6), NAVY, 1.6, 0.04)
    add_text(slide, 4.16, 1.3, 4.9, 0.26, [("2  PagerDuty cloud", 13, True, NAVY, 0)])
    add_text(slide, 4.16, 1.54, 4.9, 0.22, [("SaaS. Confirm the service region with the account owner.", 11, False, SLATE, 0)])

    add_card(slide, 4.2, 1.82, 1.55, 0.7, NAVY, NAVY, 1.0)
    add_text(slide, 4.26, 1.86, 1.44, 0.62, [("Incident", 12, True, WHITE, 0), ("Escalation policy", 11, False, WHITE, 1)])
    add_card(slide, 5.85, 1.82, 1.55, 0.7, NAVY, NAVY, 1.0)
    add_text(slide, 5.91, 1.86, 1.44, 0.62, [("Call and SMS", 12, True, WHITE, 0), ("To the on-call path", 11, False, WHITE, 1)])
    add_card(slide, 7.5, 1.82, 1.52, 0.7, WHITE, NAVY2, 1.0)
    add_text(slide, 7.56, 1.84, 1.4, 0.64, [("Sign-in", 12, True, NAVY, 0), ("Entra ID, SAML,", 10, False, SLATE, 0), ("MFA. IdP is outside.", 10, False, SLATE, 0)])

    add_card(slide, 4.2, 2.68, 4.82, 1.55, ORANGE_BG, ORANGE, 2.0)
    add_text(slide, 4.32, 2.72, 4.55, 0.4, [("PagerDuty AI Assistant", 14, True, ORANGE, 0), ("New box. Writes text back to Teams.", 11, False, SLATE, 0)])
    chips = [(4.32, "summary"), (5.5, "causes"), (6.68, "status"), (7.7, "Scribe")]
    for x, label in chips:
        add_card(slide, x, 3.55, 1.1, 0.5, WHITE, ORANGE, 1.0)
        add_text(slide, x + 0.02, 3.62, 1.06, 0.36, [(label, 11, True, NAVY, 0)], align=PP_ALIGN.CENTER)

    add_card(slide, 4.2, 4.38, 4.82, 1.15, GOLD, GOLD_LINE, 1.0)
    add_text(slide, 4.3, 4.42, 4.62, 1.05, [
        ("Cloud note", 12, True, RGBColor(0x85, 0x4D, 0x0E), 0),
        ("Transit is TLS 1.2 or higher. At rest is AES-256.", 11, False, SLATE, 2),
        ("PagerDuty states it does not train on customer data.", 11, False, SLATE, 1),
        ("Shift schedule data is not sent.", 11, False, SLATE, 1),
    ])

    # Zone 3
    add_card(slide, 10.0, 1.26, 3.15, 4.45, TEAL_BG, TEAL, 1.35, 0.04)
    add_text(slide, 10.1, 1.3, 2.95, 0.26, [("3  War room", 13, True, TEAL, 0)])
    add_card(slide, 10.14, 1.64, 2.88, 0.7, WHITE, TEAL, 1.0)
    add_text(slide, 10.22, 1.68, 2.72, 0.62, [("Microsoft Teams", 12, True, TEAL, 0), ("monitoring-2026-prod", 12, True, NAVY, 1)])
    add_card(slide, 10.14, 2.44, 2.88, 0.78, WHITE, TEAL, 1.15)
    red_dot(slide, 10.22, 2.52)
    add_text(slide, 10.4, 2.48, 2.52, 0.7, [("Primary", 14, True, TEAL, 0), ("On-call employee. Asks, then checks the reply.", 11, False, SLATE, 0)])
    add_card(slide, 10.14, 3.32, 2.88, 0.62, WHITE, GREEN, 1.0)
    add_text(slide, 10.22, 3.36, 2.72, 0.54, [("BAM-India 24x7", 12, True, GREEN, 0), ("Receives call and SMS, then escalates", 10, False, SLATE, 0)])
    add_card(slide, 10.14, 4.04, 2.88, 1.48, WHITE, TEAL, 1.0)
    add_text(slide, 10.22, 4.08, 2.72, 1.38, [
        ("In the channel", 12, True, TEAL, 0),
        ("Technology SREs", 11, False, SLATE, 2),
        ("Application Support", 11, False, SLATE, 0),
        ("Late joiner reads the reply,", 11, False, SLATE, 2),
        ("then opens DEM, Nagios,", 11, False, SLATE, 0),
        ("SCOM, or the AWS tool.", 11, False, SLATE, 0),
    ])

    # Arrows between zones
    add_arrow(slide, 3.33, 2.5, 4.05, 2.15)
    add_text(slide, 3.28, 2.15, 0.8, 0.4, [("HTTPS", 10, True, NAVY, 0), ("event", 10, True, NAVY, 0)])

    add_arrow(slide, 9.2, 2.15, 10.0, 1.95)
    add_text(slide, 9.22, 1.7, 0.75, 0.36, [("notify", 10, True, GREEN, 0)])

    add_arrow(slide, 9.2, 3.5, 10.0, 3.6)
    add_text(slide, 9.18, 3.22, 0.8, 0.28, [("call", 10, True, GREEN, 0)])

    add_arrow(slide, 10.0, 2.8, 9.2, 3.3)
    add_text(slide, 9.2, 2.55, 0.75, 0.28, [("prompt", 10, True, ORANGE, 0)])

    # Bottom records
    add_card(slide, 0.18, 5.84, 4.2, 1.05, SLATE_BG, SLATE, 1.1)
    add_text(slide, 0.28, 5.9, 4.0, 0.92, [
        ("ServiceNow / Silva", 13, True, SLATE, 0),
        ("Correlation stays here.", 12, False, SLATE, 2),
        ("Primary files it after the fact-check.", 12, False, SLATE, 0),
    ])
    add_card(slide, 4.52, 5.84, 4.2, 1.05, WHITE, NAVY2, 1.1)
    add_text(slide, 4.62, 5.9, 4.0, 0.92, [
        ("Final PIR", 13, True, NAVY, 0),
        ("Primary replaces the root-cause draft", 12, False, SLATE, 2),
        ("with confirmed evidence, then publishes.", 12, False, SLATE, 0),
    ])
    add_card(slide, 8.86, 5.84, 4.28, 1.05, SLATE_BG, RED, 1.1)
    add_text(slide, 8.96, 5.9, 4.08, 0.92, [
        ("If the assistant is down", 13, True, RED, 0),
        ("The incident is still created.", 12, False, SLATE, 2),
        ("The escalation policy still runs.", 12, False, SLATE, 0),
    ])

    add_text(slide, 0.22, 7.05, 12.9, 0.32, [
        ("Red dot is Primary. A spoken name can also appear in the Scribe note. Customer, financial, and health fields are not inputs on this picture.", 12, False, SLATE, 0),
    ])

    slide.notes_slide.notes_text_frame.text = (
        "Read left to right, the same way as the APM architecture page. "
        "Zone 1 is the detectors AXA already runs. They send an HTTPS event to PagerDuty. "
        "Zone 2 is the PagerDuty cloud. Incident, call, SMS, and escalation already live there. "
        "PagerDuty AI is the new box inside that cloud. It returns summary, a cause draft, a status draft, and a Scribe note. "
        "Sign-in is Microsoft Entra ID with SAML 2.0 and MFA. "
        "Zone 3 is the Teams channel monitoring-2026-prod. Primary asks there and checks the reply. "
        "BAM-India still gets the call and SMS. "
        "ServiceNow and Silva stay the correlation step. The PIR is published only after Primary replaces the draft cause. "
        "Confirm the PagerDuty service region with the account owner. Do not copy the APM Singapore layout onto this slide."
    )
    prs.save(OUT)


if __name__ == "__main__":
    build()
