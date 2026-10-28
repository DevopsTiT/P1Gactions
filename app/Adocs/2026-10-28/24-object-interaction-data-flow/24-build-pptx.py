#!/usr/bin/env python3
"""Object interaction map as the data-flow slide."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = (
    "/Users/k/Learnings/AIProject/CursorFiles/Daily Files/"
    "2026-10-28/24-object-interaction-data-flow/"
    "24-object-interaction-data-flow.pptx"
)

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY2 = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x0E, 0x6B, 0x5C)
TEAL_BG = RGBColor(0xE7, 0xF6, 0xF2)
BLUE_BG = RGBColor(0xE7, 0xF0, 0xF8)
ORANGE = RGBColor(0xC4, 0x56, 0x0A)
ORANGE_BG = RGBColor(0xFF, 0xF4, 0xE8)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
GREEN_BG = RGBColor(0xE8, 0xF6, 0xEE)
SLATE = RGBColor(0x37, 0x41, 0x51)
SLATE_BG = RGBColor(0xF3, 0xF4, 0xF6)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xF6, 0xCC)
GOLD_LINE = RGBColor(0xC4, 0x96, 0x0C)
BG = RGBColor(0xF7, 0xF8, 0xFB)


def set_run(paragraph, text, size, bold, color):
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    first = True
    for text, size, bold, color, before in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(before)
        p.space_after = Pt(0)
        set_run(p, text, size, bold, color)
    return box


def add_card(slide, x, y, w, h, fill, border, border_pt=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    shape.shadow.inherit = False
    return shape


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = NAVY
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()
    head = etree.SubElement(ln, qn("a:headEnd"))
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("h", "med")
    return conn


def red_dot(slide, x, y):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.16), Inches(0.16))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RED
    dot.line.fill.background()
    return dot


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Object interaction map"
    prs.core_properties.subject = "PagerDuty AI data flow, sample incident"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, 0.28, 0.06, 8.5, 0.34, [("Object interaction map", 24, True, WHITE, 0)])
    add_text(slide, 0.3, 0.38, 8.8, 0.26, [("Data flow for one critical incident. Sample ids, not a live ticket.", 13, False, RGBColor(0xD6, 0xE4, 0xF0), 0)])
    add_text(slide, 9.6, 0.16, 3.4, 0.42, [("Data Flow", 14, True, WHITE, 0), ("ALJ and ADJ (P&C)", 12, False, RGBColor(0xD6, 0xE4, 0xF0), 0)], align=PP_ALIGN.RIGHT)

    add_card(slide, 0.22, 0.82, 9.55, 0.48, GOLD, GOLD_LINE, 1.0)
    add_text(
        slide, 0.32, 0.88, 9.35, 0.38,
        [("Primary replaces the root-cause draft with confirmed evidence before the PIR is published.", 13, False, SLATE, 0)],
    )
    add_text(slide, 9.9, 0.88, 3.15, 0.38, [("Red dot = Primary, the on-call employee.", 12, False, RED, 0)])

    # 1 Dynatrace
    add_card(slide, 0.28, 1.48, 3.55, 0.62, BLUE_BG, NAVY2, 1.35)
    add_text(slide, 0.36, 1.52, 3.4, 0.54, [("Dynatrace", 11, True, NAVY2, 0), ("Problem P-88421", 14, True, NAVY, 0)])

    add_arrow(slide, 2.05, 2.1, 2.05, 2.38)

    # 2 PagerDuty + page to Primary + notify to Teams
    add_card(slide, 0.28, 2.38, 3.55, 0.7, NAVY, NAVY, 1.35)
    add_text(slide, 0.36, 2.42, 3.4, 0.62, [("PagerDuty", 11, True, RGBColor(0xF6, 0xC4, 0x53), 0), ("INC-48291", 16, True, WHITE, 0)])

    add_arrow(slide, 3.83, 2.62, 4.55, 2.62)
    add_text(slide, 3.85, 2.28, 0.7, 0.26, [("page", 11, True, NAVY, 0)], align=PP_ALIGN.CENTER)

    add_card(slide, 4.55, 2.28, 2.55, 0.72, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 4.64, 2.36)
    add_text(slide, 4.82, 2.34, 2.2, 0.6, [("Primary", 16, True, TEAL, 0), ("On-call, receives the page", 11, False, SLATE, 0)])

    add_arrow(slide, 3.83, 2.95, 7.35, 3.22)
    add_text(slide, 4.7, 2.98, 0.85, 0.24, [("notify", 11, True, GREEN, 0)])

    add_card(slide, 7.35, 2.95, 3.35, 0.7, GREEN_BG, GREEN, 1.35)
    add_text(slide, 7.45, 2.99, 3.15, 0.62, [("Teams", 11, True, GREEN, 0), ("monitoring-2026-prod", 14, True, NAVY, 0)])

    # 3 prompt
    add_arrow(slide, 2.05, 3.08, 2.05, 3.85)
    add_card(slide, 0.28, 3.85, 2.35, 0.68, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 0.36, 3.93)
    add_text(slide, 0.54, 3.92, 2.0, 0.55, [("Primary", 15, True, TEAL, 0), ("Asks in Teams", 11, False, SLATE, 0)])

    add_arrow(slide, 2.63, 4.18, 3.45, 4.18)
    add_text(slide, 2.7, 3.88, 0.75, 0.24, [("prompt", 11, True, ORANGE, 0)], align=PP_ALIGN.CENTER)

    add_card(slide, 3.45, 3.85, 3.7, 0.68, ORANGE_BG, ORANGE, 1.75)
    add_text(slide, 3.55, 3.9, 3.5, 0.58, [("PagerDuty AI Assistant", 15, True, ORANGE, 0), ("Reads the incident and the question", 11, False, SLATE, 0)])

    # 4 reply objects
    add_arrow(slide, 5.3, 4.53, 5.3, 4.82)
    add_card(slide, 0.28, 4.82, 8.55, 1.05, ORANGE_BG, ORANGE, 1.25)
    add_text(slide, 0.4, 4.86, 3.2, 0.28, [("AI reply objects", 13, True, ORANGE, 0)])
    chips = [
        (0.4, "summary", "Catch-up text"),
        (2.45, "causes", "Draft only"),
        (4.5, "status", "A person sends it"),
        (6.55, "Scribe wrap-up", "Bridge note, text"),
    ]
    for x, title, sub in chips:
        add_card(slide, x, 5.18, 1.9, 0.58, WHITE, ORANGE, 1.0)
        add_text(slide, x + 0.04, 5.2, 1.82, 0.52, [(title, 12, True, NAVY, 0), (sub, 10, False, SLATE, 0)])
    red_dot(slide, 6.62, 5.24)

    # 5 fact-check
    add_arrow(slide, 2.05, 5.87, 2.05, 6.12)
    add_card(slide, 0.28, 6.12, 2.15, 0.62, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 0.36, 6.2)
    add_text(slide, 0.54, 6.18, 1.8, 0.5, [("Primary", 15, True, TEAL, 0), ("Checks the reply", 11, False, SLATE, 0)])

    add_arrow(slide, 2.43, 6.42, 3.2, 6.42)
    add_text(slide, 2.45, 6.12, 0.85, 0.24, [("fact-check", 10, True, NAVY, 0)], align=PP_ALIGN.CENTER)

    add_card(slide, 3.2, 6.12, 4.15, 0.62, SLATE_BG, SLATE, 1.25)
    add_text(slide, 3.3, 6.14, 3.95, 0.56, [("PagerDuty actions", 11, True, SLATE, 0), ("escalate     or     resolve", 14, True, NAVY, 0)])

    # 6 PIR
    add_arrow(slide, 5.2, 6.74, 8.55, 6.95)
    add_card(slide, 8.55, 6.55, 2.15, 0.58, WHITE, NAVY2, 1.25)
    add_text(slide, 8.62, 6.58, 2.0, 0.5, [("PIR draft object", 12, True, NAVY, 0)])
    add_arrow(slide, 10.7, 6.84, 11.15, 6.84)
    add_text(slide, 10.72, 6.58, 0.45, 0.22, [("edit", 11, True, NAVY, 0)], align=PP_ALIGN.CENTER)
    add_card(slide, 11.15, 6.55, 1.95, 0.58, NAVY, NAVY, 1.25)
    add_text(slide, 11.22, 6.62, 1.8, 0.42, [("Final PIR", 14, True, WHITE, 0)])

    # verb legend
    add_card(slide, 10.85, 1.48, 2.25, 4.85, WHITE, NAVY2, 1.0)
    add_text(
        slide, 10.95, 1.55, 2.05, 4.7,
        [
            ("How to read a line", 12, True, NAVY, 0),
            ("page", 12, True, NAVY, 8),
            ("Incident reaches Primary.", 11, False, SLATE, 0),
            ("notify", 12, True, GREEN, 6),
            ("Channel monitoring-2026-prod is told.", 11, False, SLATE, 0),
            ("prompt", 12, True, ORANGE, 6),
            ("Primary asks the assistant.", 11, False, SLATE, 0),
            ("fact-check", 12, True, NAVY, 6),
            ("Primary checks, then escalates or resolves.", 11, False, SLATE, 0),
            ("edit", 12, True, NAVY, 6),
            ("Confirmed evidence replaces the draft cause.", 11, False, SLATE, 0),
        ],
    )

    notes = slide.notes_slide
    notes.notes_text_frame.text = (
        "Read top to bottom. Dynatrace problem P-88421 is on PagerDuty incident INC-48291. "
        "These ids are samples. PagerDuty pages Primary. The same incident notifies the Teams channel monitoring-2026-prod. "
        "Primary prompts PagerDuty AI Assistant. The reply objects are summary, causes, status, and the Scribe wrap-up. "
        "Causes is a draft. Primary fact-checks, then escalates or resolves in PagerDuty. "
        "Before the PIR is published, Primary replaces that root-cause draft with confirmed evidence and edits the PIR draft into the final PIR. "
        "The red dot is Primary, the on-call employee. A spoken name can also appear in the Scribe wrap-up."
    )

    prs.save(OUT)


if __name__ == "__main__":
    build()
