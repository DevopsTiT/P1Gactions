#!/usr/bin/env python3
"""Detailed data-flow pages for the Gate 0 deck."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = (
    "/Users/k/Learnings/AIProject/CursorFiles/Daily Files/"
    "2026-10-28/23-data-flow-ppt-detail/"
    "23-pagerduty-ai-data-flow-detail.pptx"
)

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY2 = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x0E, 0x6B, 0x5C)
TEAL_BG = RGBColor(0xE6, 0xF4, 0xF1)
BLUE_BG = RGBColor(0xE7, 0xF0, 0xF8)
ORANGE = RGBColor(0xC4, 0x56, 0x0A)
ORANGE_BG = RGBColor(0xFF, 0xF4, 0xE8)
SLATE = RGBColor(0x37, 0x41, 0x51)
SLATE_BG = RGBColor(0xF4, 0xF6, 0xF8)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG = RGBColor(0xFE, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD0, 0xD7, 0xE2)
BG = RGBColor(0xF7, 0xF8, 0xFB)
GOLD = RGBColor(0xF6, 0xC4, 0x53)


def set_run(paragraph, text, size, bold, color):
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.02)
    anchor = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}
    tf._txBody.bodyPr.set("anchor", anchor.get(MSO_ANCHOR.TOP, "t"))
    first = True
    for text, size, bold, color, before in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(before)
        p.space_after = Pt(0)
        set_run(p, text, size, bold, color)
    return box


def add_card(slide, x, y, w, h, fill, border, border_pt=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    shape.shadow.inherit = False
    return shape


def header(slide, title, subtitle, right_top, right_sub):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.78)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, 0.28, 0.06, 9.2, 0.38, [(title, 24, True, WHITE, 0)])
    add_text(slide, 0.3, 0.42, 9.4, 0.3, [(subtitle, 13, False, RGBColor(0xD6, 0xE4, 0xF0), 0)])
    add_text(
        slide, 10.3, 0.12, 2.75, 0.55,
        [(right_top, 13, True, WHITE, 0), (right_sub, 11, False, RGBColor(0xD6, 0xE4, 0xF0), 1)],
        align=PP_ALIGN.RIGHT,
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_hops(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "Data Flow",
        "24x7 critical-incident catch-up in Microsoft Teams.  ALJ and ADJ (P&C).",
        "Page 1 of 2",
        "After the incident exists",
    )
    add_text(
        slide, 0.28, 0.86, 12.7, 0.28,
        [("Red dot means an employee name or a Teams identity. Customer, financial, and health fields are off this page.", 13, False, SLATE, 0)],
    )

    # column labels
    add_text(slide, 0.28, 1.16, 0.4, 0.24, [("#", 11, True, NAVY2, 0)])
    add_text(slide, 0.72, 1.16, 2.5, 0.24, [("Hop", 11, True, NAVY2, 0)])
    add_text(slide, 3.55, 1.16, 5.8, 0.24, [("Data that moves", 11, True, NAVY2, 0)])
    add_text(slide, 9.85, 1.16, 3.2, 0.24, [("Personal information", 11, True, NAVY2, 0)])

    rows = [
        ("1", "Detector to PagerDuty",
         "Dynatrace, Splunk, Selenium/Jenkins, or a network tool sends the critical event. The fields are the service name, the alert text, and the problem id.",
         "No red dot on this hop.", False),
        ("2", "Page to BAM-India",
         "PagerDuty creates the incident, then places the call and the SMS under the escalation policy. BAM-India receives the page and escalates.",
         "The on-call number stays in the notification. It is not prompt text.", False),
        ("3", "PagerDuty to the assist",
         "The assist reads the id, title, priority, status, service, timeline, and the runbook linked on that service.",
         "Red dot: assignee display name. That is the on-call employee.", True),
        ("4", "Person to Teams",
         "One question in the war-room channel. The person may paste one technical line from Dynatrace or Splunk. The full log store is not connected.",
         "Red dot: Teams identity of the person who asks.", True),
        ("5", "Bridge to Scribe",
         "Scribe hears the words spoken on the Teams bridge. The result it keeps is a text note in that channel.",
         "Red dot: an employee name spoken on the bridge.", True),
        ("6", "Insights, if enabled",
         "One question, plus past PagerDuty incidents on the services this user can already see. The weekly notice is out of this pilot.",
         "Red dot: the asker’s Teams identity. History itself is operational.", True),
        ("7", "Assist back to Teams",
         "Text comes back: an SRE catch-up, a status draft, or a bridge note. A person sends any status. The tool only drafts it.",
         "Red dot if the note repeats an employee name.", True),
        ("8", "Person, then Silva",
         "A person checks the text, then acknowledges, escalates, or sends the status. That person files the ServiceNow or Silva correlation.",
         "The assist does not write this hop.", False),
    ]

    y = 1.42
    h = 0.66
    for num, hop, data, pi, marked in rows:
        fill = RED_BG if marked else WHITE
        border = RED if marked else LINE
        add_card(slide, 0.22, y, 12.9, h, fill, border, 1.15 if marked else 0.75)
        add_text(slide, 0.28, y + 0.12, 0.4, 0.4, [(num, 16, True, RED if marked else NAVY, 0)])
        add_text(slide, 0.7, y + 0.08, 2.75, 0.52, [(hop, 13, True, NAVY if not marked else RED, 0)])
        add_text(slide, 3.5, y + 0.06, 6.15, 0.56, [(data, 12, False, SLATE, 0)])
        add_text(slide, 9.75, y + 0.06, 3.25, 0.56, [(pi, 12, True if marked else False, RED if marked else SLATE, 0)])
        y += h + 0.03

    add_text(
        slide, 0.28, 7.12, 12.8, 0.3,
        [("Shift is not a hop. The on-call schedule is not sent, so Shift has no output. Customer date of birth, My Number, financial fields, and health fields are not on this flow.", 12, False, SLATE, 0)],
    )
    notes(
        slide,
        "Walk the eight hops from top to bottom. "
        "Hop 1 is the detector. Hop 2 is the call and SMS to BAM-India. The phone number stays there. "
        "Hop 3 is what the assist reads from PagerDuty. The red dot is the assignee display name. "
        "Hop 4 is the Teams question. The red dot is the asker’s identity. One technical line may be pasted. "
        "Hop 5 is bridge audio in, text note out. "
        "Hop 6 exists only if Insights is enabled. "
        "Hop 7 is the text reply. A person sends any status. "
        "Hop 8 is the person acting in PagerDuty and filing ServiceNow or Silva. "
        "Shift is not on the picture.",
    )


def slide_agents(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "Data Flow — what each agent reads and writes",
        "Same use case. Example ids below are samples, not live tickets.",
        "Page 2 of 2",
        "Teams channel only",
    )

    cards = [
        (0.22, 1.0, "SRE", TEAL, TEAL_BG,
         "Open incident: id, title, priority, status, service, timeline, escalation entries, assignee display name. The runbook linked on the service. The Teams question. One optional Dynatrace or Splunk line.",
         "A text catch-up in Teams. It repeats the incident state, then the starter checks from the runbook. A person opens the detector and compares the reply before they act."),
        (6.78, 1.0, "Scribe", NAVY2, BLUE_BG,
         "Audio of the Teams bridge. Names spoken on that bridge can be in the audio.",
         "A text note in the channel: time, what was said, and the next check. The kept result is text. Sound stays an input."),
        (0.22, 3.55, "Insights", ORANGE, ORANGE_BG,
         "Only if this agent is enabled. Past incidents on services this user can see, plus one Teams question.",
         "One sentence in Teams. It answers that question from the history the user can see. The weekly notice is out of this pilot."),
        (6.78, 3.55, "Shift", SLATE, SLATE_BG,
         "The on-call schedule, escalation-policy names, and overrides would be its input. None of that is sent.",
         "No output in this pilot. Do not draw a schedule summary on this flow."),
    ]
    for x, y, title, border, fill, inp, out in cards:
        add_card(slide, x, y, 6.32, 2.4, fill, border, 1.5)
        add_text(
            slide, x + 0.12, y + 0.08, 6.05, 2.22,
            [
                (title, 16, True, border, 0),
                ("Reads", 12, True, NAVY, 4),
                (inp, 12, False, SLATE, 1),
                ("Writes", 12, True, NAVY, 4),
                (out, 12, False, SLATE, 1),
            ],
        )

    add_card(slide, 0.22, 6.1, 12.9, 1.2, WHITE, NAVY, 1.25)
    add_text(
        slide, 0.36, 6.18, 12.6, 1.05,
        [
            ("Sample walk-through, not a live ticket", 13, True, NAVY, 0),
            ("INC-48291 is acknowledged, P2, service ebroker-api. The event names Dynatrace problem P-88421. The runbook says: check the upstream dependency, then the last deploy.", 12, False, SLATE, 2),
            ("The person asks in Teams: Catch me up on INC-48291. The reply is a catch-up. A status draft stays in the channel until a person sends it. Then the person acts in PagerDuty.", 12, False, SLATE, 1),
        ],
    )
    notes(
        slide,
        "SRE reads the open incident, the linked runbook, and the Teams question. It writes a catch-up. "
        "Scribe reads bridge audio and writes a text note. "
        "Insights, only if enabled, reads incidents the user can see and writes one sentence. "
        "Shift receives nothing and returns nothing. "
        "The sample ids INC-48291 and P-88421 are examples. "
        "A person checks the text before they acknowledge, escalate, or send a status.",
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "PagerDuty AI — Data Flow in detail"
    prs.core_properties.subject = "AXA Gate 0, ALJ and ADJ (P&C)"
    slide_hops(prs)
    slide_agents(prs)
    prs.save(OUT)


if __name__ == "__main__":
    build()
