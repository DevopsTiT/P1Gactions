#!/usr/bin/env python3
"""Proposed business architecture in the object-interaction style."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = (
    "/Users/k/Learnings/AIProject/CursorFiles/Daily Files/"
    "2026-10-28/25-architecture-object-map/"
    "25-architecture-object-map.pptx"
)

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY2 = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x0E, 0x6B, 0x5C)
TEAL_BG = RGBColor(0xE7, 0xF6, 0xF2)
BLUE_BG = RGBColor(0xE7, 0xF0, 0xF8)
ORANGE = RGBColor(0xC4, 0x56, 0x0A)
ORANGE_BG = RGBColor(0xFF, 0xF4, 0xE8)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
GREEN_BG = RGBColor(0xE8, 0xF6, 0xEE)
SLATE = RGBColor(0x37, 0x41, 0x51)
SLATE_BG = RGBColor(0xF3, 0xF4, 0xF6)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xF6, 0xCC)
GOLD_LINE = RGBColor(0xC4, 0x96, 0x0C)
BG = RGBColor(0xF7, 0xF8, 0xFB)


def set_run(paragraph, text, size, bold, color):
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    first = True
    for text, size, bold, color, before in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(before)
        p.space_after = Pt(0)
        set_run(p, text, size, bold, color)
    return box


def add_card(slide, x, y, w, h, fill, border, border_pt=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.1
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    shape.shadow.inherit = False
    return shape


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = NAVY
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()
    head = etree.SubElement(ln, qn("a:headEnd"))
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("h", "med")
    return conn


def red_dot(slide, x, y):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.14), Inches(0.14))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RED
    dot.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Proposed Business Architecture"
    prs.core_properties.subject = "Object map for the 24x7 war room"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.68))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, 0.28, 0.04, 8.6, 0.32, [("Proposed Business Architecture", 22, True, WHITE, 0)])
    add_text(slide, 0.3, 0.36, 8.8, 0.26, [("Same object style as the data-flow map. ALJ and ADJ (P&C).", 13, False, RGBColor(0xD6, 0xE4, 0xF0), 0)])
    add_text(slide, 9.8, 0.14, 3.2, 0.42, [("Architecture", 14, True, WHITE, 0), ("24x7 war room", 12, False, RGBColor(0xD6, 0xE4, 0xF0), 0)], align=PP_ALIGN.RIGHT)

    add_card(slide, 0.22, 0.8, 10.15, 0.42, GOLD, GOLD_LINE, 1.0)
    add_text(slide, 0.32, 0.86, 9.95, 0.32, [("PagerDuty AI is the only new object. Primary replaces the root-cause draft with confirmed evidence before the PIR is published.", 12, False, SLATE, 0)])
    add_text(slide, 10.5, 0.86, 2.6, 0.32, [("Red dot = Primary.", 12, True, RED, 0)])

    # Detection
    add_card(slide, 0.22, 1.38, 3.7, 1.35, BLUE_BG, NAVY2, 1.35)
    add_text(slide, 0.32, 1.42, 3.5, 1.25, [
        ("Detection zone", 14, True, NAVY, 0),
        ("Dynatrace, Splunk, Selenium/Jenkins", 11, False, SLATE, 2),
        ("VeloCloud, Nagios, Aruba UXI", 11, False, SLATE, 0),
        ("SCOM, CloudWatch, X-Ray,", 11, False, SLATE, 0),
        ("Prometheus, Kibana", 11, False, SLATE, 0),
    ])

    add_arrow(slide, 3.92, 2.0, 4.55, 2.0)
    add_text(slide, 3.85, 1.68, 0.8, 0.28, [("event", 11, True, NAVY, 0)], align=PP_ALIGN.CENTER)

    # Hub
    add_card(slide, 4.55, 1.48, 2.7, 1.15, NAVY, NAVY, 1.5)
    add_text(slide, 4.65, 1.54, 2.5, 1.02, [
        ("PagerDuty", 16, True, WHITE, 0),
        ("Incident", 12, False, WHITE, 2),
        ("Escalation policy", 12, False, WHITE, 0),
        ("Call and SMS", 12, False, WHITE, 0),
    ])

    # Three consumers
    add_arrow(slide, 7.25, 1.7, 7.85, 1.55)
    add_text(slide, 7.2, 1.32, 0.65, 0.22, [("page", 11, True, TEAL, 0)])
    add_card(slide, 7.85, 1.38, 2.55, 0.58, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 7.94, 1.46)
    add_text(slide, 8.1, 1.44, 2.2, 0.46, [("Primary", 14, True, TEAL, 0), ("On-call employee", 11, False, SLATE, 0)])

    add_arrow(slide, 7.25, 2.05, 7.85, 2.15)
    add_text(slide, 7.15, 1.82, 0.7, 0.22, [("notify", 10, True, GREEN, 0)])
    add_card(slide, 7.85, 2.05, 2.55, 0.58, GREEN_BG, GREEN, 1.25)
    add_text(slide, 7.95, 2.08, 2.35, 0.5, [("Teams", 12, True, GREEN, 0), ("monitoring-2026-prod", 11, True, NAVY, 0)])

    add_arrow(slide, 5.9, 2.63, 5.9, 2.95)
    add_text(slide, 6.0, 2.68, 0.85, 0.24, [("call / SMS", 10, True, GREEN, 0)])
    add_card(slide, 4.55, 2.95, 2.7, 0.55, GREEN_BG, GREEN, 1.25)
    add_text(slide, 4.65, 3.0, 2.5, 0.45, [("BAM-India 24x7", 13, True, GREEN, 0), ("Receives the page, then escalates", 10, False, SLATE, 0)])

    # Prompt row. Same Primary, next action.
    add_card(slide, 0.22, 3.65, 2.15, 0.62, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 0.3, 3.73)
    add_text(slide, 0.46, 3.7, 1.85, 0.52, [("Primary", 14, True, TEAL, 0), ("Asks in the channel", 11, False, SLATE, 0)])
    add_arrow(slide, 2.37, 3.96, 3.05, 3.96)
    add_text(slide, 2.4, 3.68, 0.65, 0.24, [("prompt", 11, True, ORANGE, 0)], align=PP_ALIGN.CENTER)
    add_card(slide, 3.05, 3.65, 3.35, 0.62, ORANGE_BG, ORANGE, 2.0)
    add_text(slide, 3.15, 3.68, 3.15, 0.55, [("PagerDuty AI Assistant", 13, True, ORANGE, 0), ("The only new object", 11, False, SLATE, 0)])

    # Reply objects
    add_arrow(slide, 4.7, 4.27, 4.7, 4.52)
    add_card(slide, 0.22, 4.52, 7.55, 0.95, ORANGE_BG, ORANGE, 1.15)
    add_text(slide, 0.32, 4.54, 3.4, 0.26, [("Reply objects in Teams", 12, True, ORANGE, 0)])
    chips = [
        (0.34, "summary"),
        (2.15, "causes"),
        (3.96, "status"),
        (5.77, "Scribe wrap-up"),
    ]
    for x, label in chips:
        add_card(slide, x, 4.86, 1.7, 0.48, WHITE, ORANGE, 1.0)
        add_text(slide, x + 0.04, 4.92, 1.62, 0.36, [(label, 12, True, NAVY, 0)], align=PP_ALIGN.CENTER)
    red_dot(slide, 5.84, 4.92)

    # Bottom actions
    add_arrow(slide, 1.2, 5.47, 1.2, 5.75)
    add_card(slide, 0.22, 5.75, 1.9, 0.7, TEAL_BG, TEAL, 1.5)
    red_dot(slide, 0.3, 5.83)
    add_text(slide, 0.46, 5.82, 1.6, 0.55, [("Primary", 14, True, TEAL, 0), ("Fact-check", 11, False, SLATE, 0)])

    add_arrow(slide, 2.12, 6.05, 2.55, 6.05)
    add_card(slide, 2.55, 5.75, 2.35, 0.7, SLATE_BG, SLATE, 1.15)
    add_text(slide, 2.65, 5.8, 2.15, 0.6, [("PagerDuty actions", 11, True, SLATE, 0), ("escalate or resolve", 12, True, NAVY, 0)])

    add_arrow(slide, 4.9, 6.05, 5.25, 6.05)
    add_text(slide, 4.9, 5.75, 0.4, 0.22, [("edit", 10, True, NAVY, 0)])
    add_card(slide, 5.25, 5.75, 1.7, 0.7, WHITE, NAVY2, 1.15)
    add_text(slide, 5.32, 5.86, 1.55, 0.5, [("Final PIR", 13, True, NAVY, 0), ("After the draft cause is replaced", 10, False, SLATE, 0)])

    add_arrow(slide, 6.95, 6.1, 7.3, 6.35)
    add_text(slide, 6.95, 5.78, 0.45, 0.22, [("file", 10, True, SLATE, 0)])
    add_card(slide, 7.3, 6.15, 2.55, 0.62, SLATE_BG, SLATE, 1.15)
    add_text(slide, 7.4, 6.2, 2.35, 0.52, [("ServiceNow / Silva", 13, True, SLATE, 0), ("Correlation after Primary acts", 10, False, SLATE, 0)])

    # Side note for late joiners
    add_card(slide, 10.55, 1.38, 2.55, 4.15, WHITE, NAVY2, 1.0)
    add_text(slide, 10.65, 1.46, 2.35, 3.95, [
        ("Who owns a box", 13, True, NAVY, 0),
        ("Detection", 12, True, NAVY2, 8),
        ("The team that already owns that tool.", 11, False, SLATE, 0),
        ("PagerDuty", 12, True, NAVY2, 6),
        ("Technology DevOps.", 11, False, SLATE, 0),
        ("Teams channel", 12, True, GREEN, 6),
        ("Technology SREs and Application Support.", 11, False, SLATE, 0),
        ("Late joiner", 12, True, TEAL, 6),
        ("Reads the reply, then opens one tool.", 11, False, SLATE, 0),
        ("DBA uses DEM.", 11, False, SLATE, 2),
        ("Network uses Nagios.", 11, False, SLATE, 0),
        ("Axa GO uses SCOM.", 11, False, SLATE, 0),
        ("Cloud Broker uses AWS tools.", 11, False, SLATE, 0),
    ])

    add_text(slide, 0.22, 7.12, 10.1, 0.28, [("Shift is not an object on this picture. If the assistant is down, the incident and the escalation policy still run.", 12, False, SLATE, 0)])

    slide.notes_slide.notes_text_frame.text = (
        "This is the standing architecture, drawn like the object interaction map. "
        "Detection sends the critical event to PagerDuty. "
        "PagerDuty pages Primary, notifies the Teams channel monitoring-2026-prod, and calls BAM-India. "
        "Primary prompts PagerDuty AI. That assistant is the only new object. "
        "The reply in Teams has four objects: summary, causes, status, and the Scribe wrap-up. "
        "Causes stays a draft until Primary replaces it with confirmed evidence. "
        "Primary then escalates or resolves, edits the final PIR, and files ServiceNow or Silva. "
        "A late joiner reads the same reply and opens the one tool their team owns."
    )
    prs.save(OUT)


if __name__ == "__main__":
    build()
