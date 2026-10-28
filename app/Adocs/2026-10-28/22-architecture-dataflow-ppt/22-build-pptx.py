#!/usr/bin/env python3
"""Two-slide Gate 0 deck: architecture zones and data flow."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = (
    "/Users/k/Learnings/AIProject/CursorFiles/Daily Files/"
    "2026-10-28/22-architecture-dataflow-ppt/"
    "22-pagerduty-ai-architecture-data-flow.pptx"
)

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY2 = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x0E, 0x6B, 0x5C)
TEAL_BG = RGBColor(0xE6, 0xF4, 0xF1)
BLUE_BG = RGBColor(0xE7, 0xF0, 0xF8)
GREEN_BG = RGBColor(0xE8, 0xF6, 0xEE)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
ORANGE = RGBColor(0xC4, 0x56, 0x0A)
ORANGE_BG = RGBColor(0xFF, 0xF4, 0xE8)
YELLOW = RGBColor(0xFF, 0xF6, 0xCC)
YELLOW_LINE = RGBColor(0xC4, 0x96, 0x0C)
SLATE = RGBColor(0x37, 0x41, 0x51)
SLATE_BG = RGBColor(0xF3, 0xF4, 0xF6)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG = RGBColor(0xFE, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x94, 0xA3, 0xB8)
BG = RGBColor(0xF8, 0xFA, 0xFC)


def set_run(paragraph, text, size, bold, color, font="Calibri"):
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.paragraphs[0].alignment = align
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}.get(anchor, "t"))
    except Exception:
        pass
    first = True
    for text, size, bold, color, space_before in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(space_before)
        p.space_after = Pt(0)
        set_run(p, text, size, bold, color)
    return box


def add_card(slide, x, y, w, h, fill, border, border_pt=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    shape.shadow.inherit = False
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=NAVY):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(1.75)
    ln = conn.line._get_or_add_ln()
    head = etree.SubElement(ln, qn("a:headEnd"))
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("h", "med")
    return conn


def add_notes(slide, text):
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "PagerDuty AI — Architecture and Data Flow"
    prs.core_properties.subject = "AXA Gate 0, ALJ and ADJ (P&C), 24x7 war room"
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    architecture(s1)
    s2 = prs.slides.add_slide(blank)
    data_flow(s2)
    prs.save(OUT)


def architecture(slide):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.92)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    add_text(
        slide, 0.32, 0.08, 10.2, 0.42,
        [("Proposed Business Architecture", 26, True, WHITE, 0)],
    )
    add_text(
        slide, 0.34, 0.48, 10.4, 0.36,
        [("24x7 war-room assist for ALJ and ADJ (P&C). The critical event already exists.", 13, False, RGBColor(0xD6, 0xE4, 0xF0), 0)],
    )
    add_text(
        slide, 10.6, 0.22, 2.4, 0.5,
        [("Gate 0", 14, True, WHITE, 0), ("Document 3.1", 11, False, RGBColor(0xD6, 0xE4, 0xF0), 0)],
        align=PP_ALIGN.RIGHT,
    )

    add_card(slide, 0.22, 1.12, 3.25, 5.6, BLUE_BG, NAVY2, 1.5)
    add_text(
        slide, 0.36, 1.24, 2.98, 5.3,
        [
            ("DETECTION ZONE", 14, True, NAVY, 0),
            ("Providers of the signal", 12, False, TEAL, 2),
            ("Dynatrace", 13, False, SLATE, 8),
            ("Splunk", 13, False, SLATE, 1),
            ("Selenium / Jenkins", 13, False, SLATE, 1),
            ("VeloCloud, Nagios, Aruba UXI", 13, False, SLATE, 1),
            ("SCOM", 13, False, SLATE, 1),
            ("CloudWatch, X-Ray,", 13, False, SLATE, 1),
            ("Prometheus, Kibana", 13, False, SLATE, 0),
            ("Technology DevOps owns Dynatrace,", 11, False, NAVY2, 10),
            ("Splunk, Selenium/Jenkins, and PagerDuty.", 11, False, NAVY2, 0),
            ("DBA, Network, Axa GO, and Cloud Broker", 11, False, NAVY2, 2),
            ("own the specialist tools.", 11, False, NAVY2, 0),
        ],
    )

    add_card(slide, 4.72, 2.05, 2.5, 2.7, NAVY, NAVY, 1.5)
    add_text(
        slide, 4.84, 2.16, 2.28, 2.45,
        [
            ("PAGERDUTY", 15, True, WHITE, 0),
            ("Hub", 12, False, RGBColor(0xF6, 0xC4, 0x53), 2),
            ("Incident", 13, False, WHITE, 8),
            ("Escalation policy", 13, False, WHITE, 2),
            ("Call and SMS", 13, False, WHITE, 2),
            ("Owner: Technology DevOps", 11, False, RGBColor(0xD6, 0xE4, 0xF0), 8),
        ],
    )

    add_card(slide, 8.55, 1.12, 2.15, 1.75, GREEN_BG, GREEN, 1.5)
    add_text(
        slide, 8.66, 1.2, 1.95, 1.58,
        [
            ("BAM-INDIA", 14, True, GREEN, 0),
            ("Consumer of the page", 11, False, SLATE, 2),
            ("Receives the call and SMS", 12, False, SLATE, 4),
            ("Escalates into the war room", 12, False, SLATE, 1),
        ],
    )

    add_card(slide, 10.85, 1.12, 2.25, 1.75, YELLOW, YELLOW_LINE, 1.5)
    add_text(
        slide, 10.95, 1.18, 2.05, 1.6,
        [
            ("YELLOW NOTE", 11, True, RGBColor(0x85, 0x4D, 0x0E), 0),
            ("Only new box: PagerDuty AI.", 11, False, SLATE, 3),
            ("It sits in the Teams channel.", 11, False, SLATE, 1),
            ("Shift is not drawn.", 11, False, SLATE, 1),
            ("Assist down: paging still runs.", 11, False, SLATE, 1),
        ],
    )

    add_card(slide, 8.55, 3.08, 4.55, 3.15, TEAL_BG, TEAL, 1.5)
    add_text(
        slide, 8.7, 3.16, 4.25, 1.2,
        [
            ("WAR ROOM", 14, True, TEAL, 0),
            ("Consumers of the briefing", 11, False, SLATE, 1),
            ("Technology SREs and Application Support", 12, False, SLATE, 3),
            ("Microsoft Teams channel", 12, False, SLATE, 0),
        ],
    )
    add_card(slide, 8.72, 4.55, 2.15, 1.52, ORANGE_BG, ORANGE, 2.25)
    add_text(
        slide, 8.82, 4.58, 1.98, 1.42,
        [
            ("PAGERDUTY AI", 12, True, ORANGE, 0),
            ("The new box", 10, True, ORANGE, 0),
            ("SRE catch-up", 11, False, SLATE, 2),
            ("Status draft", 11, False, SLATE, 0),
            ("Scribe bridge note", 11, False, SLATE, 0),
            ("Insights, if enabled", 11, False, SLATE, 0),
        ],
    )
    add_text(
        slide, 11.0, 4.55, 1.95, 1.55,
        [
            ("Late joiner", 12, True, TEAL, 0),
            ("DBA → DEM", 11, False, SLATE, 3),
            ("Network → Nagios", 11, False, SLATE, 1),
            ("Axa GO → SCOM", 11, False, SLATE, 1),
            ("Cloud Broker → AWS", 11, False, SLATE, 1),
        ],
    )

    add_card(slide, 4.72, 5.2, 2.5, 1.45, SLATE_BG, SLATE, 1.5)
    add_text(
        slide, 4.84, 5.28, 2.28, 1.28,
        [
            ("SERVICENOW / SILVA", 13, True, SLATE, 0),
            ("Correlation", 12, False, TEAL, 2),
            ("Filed after a person has", 12, False, SLATE, 4),
            ("checked the text and acted.", 12, False, SLATE, 0),
        ],
    )

    add_arrow(slide, 3.47, 3.3, 4.72, 3.3)
    add_text(slide, 3.4, 2.95, 1.3, 0.28, [("critical event", 10, True, NAVY, 0)], align=PP_ALIGN.CENTER)

    add_arrow(slide, 7.22, 2.35, 8.55, 1.85)
    add_text(slide, 7.22, 1.48, 1.25, 0.28, [("call / SMS", 10, True, GREEN, 0)], align=PP_ALIGN.CENTER)

    add_arrow(slide, 7.22, 3.85, 8.55, 3.85)
    add_text(slide, 7.2, 3.5, 1.3, 0.28, [("escalation", 10, True, TEAL, 0)], align=PP_ALIGN.CENTER)

    add_arrow(slide, 5.5, 4.75, 5.5, 5.2)
    add_text(slide, 5.65, 4.82, 1.4, 0.28, [("correlation", 10, True, SLATE, 0)])

    add_text(
        slide, 0.22, 6.9, 12.8, 0.4,
        [("Left box sends the signal. Center hub pages BAM-India. Orange box is the only addition. Employee names are marked on the next slide.", 12, False, SLATE, 0)],
    )

    add_notes(
        slide,
        "This slide is the business picture for the 24x7 war room. "
        "Read it like the EIP page: zones as boxes, PagerDuty as the hub, a label on every line. "
        "The detection zone provides the signal. Each tool keeps its current owner. "
        "PagerDuty already creates the incident, places the call and SMS, and runs the escalation policy. "
        "BAM-India consumes the page and escalates. "
        "Technology SREs and Application Support consume the briefing in Teams. "
        "PagerDuty AI is the thicker orange box. It writes a catch-up, a status draft, or a bridge note. "
        "A person checks the text, then acts in PagerDuty. "
        "ServiceNow and Silva are the correlation after that. "
        "Shift is not drawn. If the assist is down, the incident and the escalation policy still run. "
        "This uses the EIP layout. The EIP mediators stay on the EIP page.",
    )


def data_flow(slide):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.92)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    add_text(slide, 0.32, 0.08, 9.5, 0.42, [("Data Flow", 26, True, WHITE, 0)])
    add_text(
        slide, 0.34, 0.48, 10.2, 0.36,
        [("Use case: 24x7 critical-incident catch-up in Microsoft Teams", 13, False, RGBColor(0xD6, 0xE4, 0xF0), 0)],
    )
    add_text(
        slide, 10.5, 0.22, 2.5, 0.5,
        [("ALJ and ADJ", 14, True, WHITE, 0), ("P&C", 11, False, RGBColor(0xD6, 0xE4, 0xF0), 0)],
        align=PP_ALIGN.RIGHT,
    )

    # legend
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.36), Inches(1.08), Inches(0.22), Inches(0.22))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RED
    dot.line.fill.background()
    add_text(
        slide, 0.64, 1.02, 12.3, 0.34,
        [("Red dot = personal information: an employee display name, a Teams identity, or a name spoken on the bridge.", 13, False, SLATE, 0)],
    )

    steps = [
        (0.22, "1", "Detector", "Service name\nAlert text\nProblem id", False, BLUE_BG, NAVY2),
        (2.38, "2", "PagerDuty", "Incident fields\nRunbook link\nRed dot: assignee", True, NAVY, WHITE),
        (4.54, "3", "Teams", "The question\nOne technical line\nRed dot: who asks", True, TEAL_BG, TEAL),
        (6.70, "4", "PagerDuty AI", "SRE catch-up\nScribe note\nInsights if enabled", False, ORANGE_BG, ORANGE),
        (8.86, "5", "Text reply", "Catch-up or note\nStatus draft\nRed dot if name repeats", True, RED_BG, RED),
        (11.02, "6", "Person", "Checks the text\nActs in PagerDuty\nFiles correlation", False, SLATE_BG, SLATE),
    ]

    for x, num, title, body, pi, fill, border in steps:
        title_color = WHITE if fill == NAVY else border
        body_color = WHITE if fill == NAVY else SLATE
        add_card(slide, x, 1.52, 1.95, 2.55, fill, border, 1.75 if pi else 1.25)
        lines = [(num + "  " + title, 12, True, title_color, 0)]
        parts = body.split("\n")
        for i, part in enumerate(parts):
            mark = part.startswith("Red dot")
            if mark and fill == NAVY:
                color = RGBColor(0xF6, 0xC4, 0x53)
            elif mark:
                color = RED
            else:
                color = body_color
            lines.append((part, 12, mark, color, 4 if i == 0 else 1))
        add_text(slide, x + 0.08, 1.6, 1.8, 2.35, lines)

    centers_y = 2.75
    for x, *_rest in steps[:-1]:
        add_arrow(slide, x + 1.95, centers_y, x + 2.16, centers_y)

    add_text(
        slide, 0.28, 4.2, 12.7, 0.32,
        [("Call and SMS stay inside PagerDuty and go to BAM-India. The on-call number is not copied into the AI prompt. Shift has no arrow.", 13, False, SLATE, 0)],
    )

    # three red-dot cards
    cards = [
        ("Red dot on step 2", "Assignee display name on the PagerDuty incident. That is the on-call employee’s name."),
        ("Red dot on step 3", "Teams identity of the person who types the question in the war-room channel."),
        ("Red dot on step 5", "An employee name spoken on the bridge, if the Scribe note repeats it."),
    ]
    for i, (title, body) in enumerate(cards):
        x = 0.28 + i * 4.35
        add_card(slide, x, 4.65, 4.15, 1.45, RED_BG, RED, 1.25)
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12), Inches(4.78), Inches(0.18), Inches(0.18))
        oval.fill.solid()
        oval.fill.fore_color.rgb = RED
        oval.line.fill.background()
        add_text(
            slide, x + 0.38, 4.72, 3.65, 1.25,
            [
                (title, 13, True, RED, 0),
                (body, 12, False, SLATE, 4),
            ],
        )

    add_text(
        slide, 0.28, 6.25, 12.7, 0.95,
        [
            ("Left off this picture", 13, True, NAVY, 0),
            ("Customer date of birth, My Number, financial fields, and health fields.", 13, False, SLATE, 2),
            ("The paging phone number stays in the PagerDuty notification.", 13, False, SLATE, 0),
            ("The on-call schedule is not sent, so Shift produces no output in this pilot.", 13, False, SLATE, 0),
        ],
    )

    add_notes(
        slide,
        "Same path as the architecture slide, drawn as data. "
        "Use case: 24x7 critical-incident catch-up in Microsoft Teams, for ALJ and ADJ (P&C), after PagerDuty has opened the incident. "
        "Step 1 is the detector: service name, alert text, problem id. "
        "Step 2 is the incident. Red dot is the assignee display name. The on-call phone stays in the notification and is not prompt text. "
        "Step 3 is the Teams question. Red dot is the asker’s Teams identity. At most one technical line is pasted. "
        "Step 4 is PagerDuty AI: SRE catch-up, Scribe note from bridge audio, Insights only if enabled. "
        "Step 5 is text back in Teams. Red dot if a spoken name is repeated. "
        "Step 6: a person checks the text, acts in PagerDuty, and files ServiceNow / Silva. The assist does not write that hop. "
        "Customer date of birth, My Number, financial fields, and health fields are not on this picture.",
    )


if __name__ == "__main__":
    build()
